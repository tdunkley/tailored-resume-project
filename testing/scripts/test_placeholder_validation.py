import json
from docx import Document

# Function to extract placeholders from the Word template
def extract_placeholders_from_template(template_path):
    doc = Document(template_path)
    placeholders = set()

    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "{{" in text and "}}" in text:
            placeholders.update(
                placeholder.strip() for placeholder in text.split() if "{{" in placeholder and "}}" in placeholder
            )

    return placeholders

# Function to validate placeholders against JSON keys
def validate_placeholders(template_placeholders, json_keys):
    stripped_placeholders = {placeholder.strip("{}") for placeholder in template_placeholders}
    missing_placeholders = stripped_placeholders - json_keys
    unused_keys = json_keys - stripped_placeholders

    return missing_placeholders, unused_keys


# Test script
template_path = "templates/resume_template.docx"  # Path to your template file
json_path = "data/tailored_resume.json"  # Path to your JSON file

# Extract placeholders from the template
template_placeholders = extract_placeholders_from_template(template_path)

# Load JSON keys
with open(json_path, "r") as file:
    json_data = json.load(file)
json_keys = set(json_data.keys())

# Validate
missing_placeholders, unused_keys = validate_placeholders(template_placeholders, json_keys)

print("Validation Results:")
if missing_placeholders:
    print(f"Error: These placeholders are missing in the JSON: {missing_placeholders}")
else:
    print("No missing placeholders in the JSON.")

if unused_keys:
    print(f"Warning: These keys are in the JSON but not used in the template: {unused_keys}")
else:
    print("No unused keys in the JSON.")
