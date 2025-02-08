from docx import Document

# Function to extract placeholders from the Word template
def extract_placeholders_from_template(template_path):
    doc = Document(template_path)
    placeholders = set()

    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "{{" in text and "}}" in text:
            placeholders.update(
                placeholder.strip() for placeholder in text.split() if "{{" in placeholder and "}}" in placeholder
            )

    return placeholders

# Test the function
template_path = "templates/resume_template.docx"  # Path to your template file
placeholders = extract_placeholders_from_template(template_path)

print("Extracted placeholders from template:")
print(placeholders)
